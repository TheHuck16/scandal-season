#!/usr/bin/env python3
"""Build Season 8/9/10 complete review DOCX volumes from chapter .md files.

Usage: python3 build_review.py 8 "The Reckoning"
Reads Content/story/season-{eight,nine,ten}-chapter-*.md (+ pilot file),
produces ~/workspace/your_files/Season {N} — {Title} — Complete Review Draft.docx
"""
import os
import re
import sys
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH

STORY = os.path.expanduser("~/workspace/scandal-season/Content/story")
OUT = os.path.expanduser("~/workspace/your_files")

SEASON_WORDS = {8: "eight", 9: "nine", 10: "ten"}

def chapter_files(season):
    w = SEASON_WORDS[season]
    files = []
    pilot = os.path.join(STORY, f"season-{w}-chapters-1-3-pilot.md")
    if os.path.exists(pilot):
        files.append((1, pilot))
    for c in range(4, 31):
        p = os.path.join(STORY, f"season-{w}-chapter-{c:02d}.md")
        if os.path.exists(p):
            files.append((c, p))
    return files

def clean_md(text):
    """Strip markdown formatting for DOCX paragraphs."""
    text = re.sub(r'\*\*(.+?)\*\*', r'\1', text)
    text = re.sub(r'(?<!\*)\*(?!\*)(.+?)(?<!\*)\*(?!\*)', r'\1', text)
    text = text.replace('\\*', '*')
    return text.strip()

def build(season, title):
    doc = Document()
    style = doc.styles['Normal']
    style.font.name = 'Georgia'
    style.font.size = Pt(11)

    # Title page
    for _ in range(4):
        doc.add_paragraph()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(f"Season {title.split('—')[0].strip()}" if '—' in title else f"Season {season}")
    r.font.size = Pt(28)
    r.bold = True
    # Simpler: use the passed title
    doc.paragraphs[-1].clear()
    t = doc.add_paragraph()
    t.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t.add_run(f"Season {season} — {title}")
    r.font.size = Pt(28)
    r.bold = True
    t2 = doc.add_paragraph()
    t2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t2.add_run("Complete Review Draft")
    r.font.size = Pt(18)
    r.italic = True
    t3 = doc.add_paragraph()
    t3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = t3.add_run("STATUS: FULL DRAFT — awaiting Beth review.")
    r.font.size = Pt(12)
    doc.add_page_break()

    files = chapter_files(season)
    assert len(files) == 28, f"expected 28 files (pilot + 27), got {len(files)}"

    # Track chapter numbering: pilot covers ch1-3, then ch4-30
    for ch_start, path in files:
        text = open(path).read()
        scenes = re.split(r'^(### L\d+\.S\d+ .+)$', text, flags=re.M)

        # Chapter headings
        if 'pilot' in path:
            # Pilot has ## Chapter N — "Title" boundaries
            # Split body by chapter boundaries and emit headings
            ch_bounds = [(m.start(), m.group(1), clean_md(m.group(2)))
                         for m in re.finditer(r'^## Chapter (\d+)\s*[—–-]\s*(.+)$', text, re.M)]
            # We'll emit chapter headings inline as we walk scenes
            pending_ch = {i: (num, ttl) for i, (pos, num, ttl) in enumerate(ch_bounds)}
            # Map scene index to chapter: find which chapter each scene belongs to
            scene_positions = [m.start() for m in re.finditer(r'^### L\d+\.S\d+ ', text, re.M)]
            # Build list of (scene_idx, chapter_num, chapter_title)
            ch_for_scene = []
            for sp in scene_positions:
                cur = ch_bounds[0]
                for b in ch_bounds:
                    if b[0] <= sp:
                        cur = b
                ch_for_scene.append((cur[1], cur[2]))
        else:
            m = re.search(r'^# Scandal Season — Season \w+, Chapter (\d+):\s*"([^"]+)"', text, re.M)
            if m:
                doc.add_heading(f"Chapter {m.group(1)} — {m.group(2)}", level=1)
            else:
                m2 = re.search(r'chapter-(\d+)', path)
                doc.add_heading(f"Chapter {m2.group(1)}", level=1)
            ch_for_scene = None

        # Walk scenes
        last_ch = None
        for i in range(1, len(scenes), 2):
            header = scenes[i].strip()
            body = scenes[i+1] if i+1 < len(scenes) else ""
            if ch_for_scene is not None:
                scene_idx = (i - 1) // 2
                num, ttl = ch_for_scene[scene_idx]
                if num != last_ch:
                    doc.add_heading(f"Chapter {num} — {ttl}", level=1)
                    last_ch = num
            # Scene header as Heading 2
            hdr_clean = re.sub(r'^### ', '', header)
            doc.add_heading(clean_md(hdr_clean), level=2)
            # Body: blockquotes and paragraphs
            for line in body.split('\n'):
                line = line.strip()
                if not line:
                    continue
                if line.startswith('>'):
                    p = doc.add_paragraph(clean_md(line.lstrip('> ').strip()))
                    p.paragraph_format.left_indent = Pt(18)
                    p.paragraph_format.space_after = Pt(4)
                elif line.startswith('*Purpose:') or line.startswith('*Animation:'):
                    p = doc.add_paragraph()
                    r = p.add_run(clean_md(line))
                    r.italic = True
                    r.font.size = Pt(10)
                    r.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                elif line.startswith('*★') or line.startswith('★'):
                    p = doc.add_paragraph()
                    r = p.add_run(clean_md(line))
                    r.bold = True
                elif line.startswith('- **') or line.startswith('  - '):
                    p = doc.add_paragraph(clean_md(line.lstrip('- ').strip()), style='List Bullet')
                elif line.startswith('#'):
                    continue
                else:
                    doc.add_paragraph(clean_md(line))

    out = os.path.join(OUT, f"Season {season} — {title} — Complete Review Draft.docx")
    doc.save(out)
    print(f"saved: {out}")
    # Verify
    d2 = Document(out)
    h1 = sum(1 for p in d2.paragraphs if p.style.name == 'Heading 1')
    h2 = sum(1 for p in d2.paragraphs if p.style.name == 'Heading 2')
    print(f"Heading1: {h1}, Heading2 (scenes): {h2}")

if __name__ == '__main__':
    build(int(sys.argv[1]), sys.argv[2])
